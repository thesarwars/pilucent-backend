"""The parse contract (attendance intake backend spec §4): bytes and a
device profile in, punches and errors out. Pure -- no database, no clock.

The profile is data; nothing here knows a brand name (§1 rule 7). The date
format comes from the profile and is never guessed per row: 03/04/2027 is 3
April or 4 March depending only on the profile (§2.2).
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone as dt_timezone
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

DHAKA = ZoneInfo("Asia/Dhaka")

FORMATS = ("CSV", "TSV", "FIXED", "XLSX", "CUSTOM")
ENCODINGS = {"UTF-8": "utf-8-sig", "UTF-16": "utf-16", "Windows-1252": "cp1252"}
ROW_STRUCTURES = ("EVENT_LOG", "PAIRED_COLUMNS")
DIRECTIONS = ("EXPLICIT_COLUMN", "ALTERNATE", "FIRST_LAST", "PAIRED")
MATCHING = ("DEVICE_ID_MAP", "EMPLOYEE_CODE", "NID")
DATE_FORMATS = ("DD/MM/YYYY", "YYYY-MM-DD", "DD-MM-YY", "MM/DD/YYYY")
TIME_FORMATS = ("HH:mm", "HH:mm:ss", "hh:mm A")
TARGETS = ("deviceUserId", "employeeCode", "date", "time", "inTime", "outTime", "direction", "terminal",
           "datetime")
OUT_WORDS = re.compile(r"OUT|1|EXIT")
MAX_BYTES = 25 * 1024 * 1024


@dataclass
class Profile:
    """What the parser reads from a device profile."""

    format: str = "CSV"
    encoding: str = "UTF-8"
    delimiter: str = ","
    header_rows: int = 1
    row_structure: str = "EVENT_LOG"
    date_format: str = "DD/MM/YYYY"
    time_format: str = "HH:mm:ss"
    combined_datetime: bool = False
    direction_handling: str = "EXPLICIT_COLUMN"
    employee_matching: str = "DEVICE_ID_MAP"
    timezone: str = "Asia/Dhaka"
    column_map: list = field(default_factory=list)  # [{source, target, start?, width?}]
    pairs: list | None = None  # [{in, out}] by source column, PAIRED_COLUMNS
    century_base: int = 2000  # DD-MM-YY: '27' is century_base + 27


@dataclass
class Punch:
    row: int
    raw_row: str
    device_user_id: str
    punched_at: datetime  # aware, UTC
    punch_date: date  # Dhaka local
    punch_time: time  # Dhaka local
    direction: str | None
    direction_source: str
    terminal: str = ""
    discarded: bool = False  # FIRST_LAST middles: stored as skipped, never dropped


@dataclass
class Result:
    punches: list
    errors: list
    rows_total: int
    rows_parsed: int
    date_range: tuple | None
    device_ids: dict  # id -> punch count
    discarded: int = 0


class ParseRefused(Exception):
    """The file cannot be read under this profile at all (not a row error)."""


# ------------------------------------------------------------------ values


def parse_date(value, fmt, century_base=2000):
    """A date in exactly the profile's format, or None. 31/02 is caught."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value or "").strip()
    parts = re.split(r"[/\-.]", s)
    if len(parts) != 3 or not all(p.isdigit() for p in parts):
        return None
    if fmt == "YYYY-MM-DD":
        y, m, d = parts
    elif fmt == "MM/DD/YYYY":
        m, d, y = parts
    else:  # DD/MM/YYYY, DD-MM-YY
        d, m, y = parts
    if fmt == "DD-MM-YY":
        if len(y) != 2:
            return None
        y = century_base + int(y)
    elif len(str(y)) != 4:
        return None
    try:
        return date(int(y), int(m), int(d))
    except ValueError:
        return None


TIME_RE = re.compile(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?\s*([AaPp][Mm])?$")


def parse_time(value, fmt):
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    m = TIME_RE.match(str(value or "").strip())
    if not m:
        return None
    h, mi, sec, ampm = int(m.group(1)), int(m.group(2)), int(m.group(3) or 0), m.group(4)
    if fmt == "hh:mm A":
        if not ampm or not 1 <= h <= 12:
            return None
        h = (0 if h == 12 else h) + (12 if ampm.upper() == "PM" else 0)
    elif ampm:
        return None
    if fmt == "HH:mm:ss" and m.group(3) is None:
        return None
    if h > 23 or mi > 59 or sec > 59:
        return None
    return time(h, mi, sec)


def zone(name):
    try:
        return ZoneInfo(name or "Asia/Dhaka")
    except (ZoneInfoNotFoundError, ValueError):
        raise ParseRefused(f"Unknown time zone {name!r}.") from None


def localise(d, t, tz):
    """The device's wall time in its zone -> the instant, and the Dhaka date
    and time the register keys by (§2.1). A device set to UTC puts an 18:30
    punch on the next Dhaka day; that is the point of storing both."""
    instant = datetime.combine(d, t).replace(tzinfo=tz)
    local = instant.astimezone(DHAKA)
    return instant.astimezone(dt_timezone.utc), local.date(), local.time()


# ------------------------------------------------------------------- rows


def split_line(line, delimiter):
    """RFC 4180, doubled quotes (§4.4): a terminal name with a comma is one field."""
    return [c.strip() for c in next(csv.reader([line], delimiter=delimiter, quotechar='"', doublequote=True,
                                               skipinitialspace=False), [])]


def decode(raw, encoding):
    codec = ENCODINGS.get(encoding)
    if codec is None:
        raise ParseRefused(f"Unknown encoding {encoding!r}.")
    try:
        return raw.decode(codec)
    except UnicodeDecodeError as exc:
        raise ParseRefused(f"The file is not {encoding}: byte {exc.start} cannot be read. Check the profile's "
                           "encoding.") from None


def _xlsx_rows(raw):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 -- any unreadable workbook is the same answer
        raise ParseRefused("The file is not a readable XLSX workbook.") from None
    ws = wb.worksheets[0]
    for row in ws.iter_rows(values_only=True):
        cells = ["" if v is None else v for v in row]
        if any(str(c).strip() for c in cells):
            yield cells, "\t".join(str(c) for c in cells)


def rows_of(raw, p):
    """(cells, raw_text) for every non-blank line, headers included."""
    if p.format == "XLSX":
        yield from _xlsx_rows(raw)
        return
    if p.format == "CUSTOM":
        raise ParseRefused("A CUSTOM format needs its own parser; none is installed for this profile.")
    text = decode(raw, p.encoding)
    for line in text.splitlines():
        if not line.strip():
            continue
        if p.format == "FIXED":
            yield line, line.rstrip("\r\n")
        else:
            yield split_line(line, "\t" if p.format == "TSV" else (p.delimiter or ",")), line


def _indexer(p, header):
    """target -> how to read it from a row."""
    out = {}
    for c in p.column_map:
        target = c.get("target")
        if target not in TARGETS:
            continue
        if p.format == "FIXED":
            start, width = int(c.get("start") or 0), int(c.get("width") or 0)
            out[target] = (lambda line, s=start, w=width: line[s:s + w].strip()) if width else None
        else:
            names = [str(h).strip() for h in header]
            src = str(c.get("source") or "").strip()
            i = names.index(src) if src in names else None
            out[target] = (lambda cells, i=i: cells[i] if i is not None and i < len(cells) else "") \
                if i is not None else None
    return out


def _pairs(p, header):
    """(field, direction, reader) for each declared in/out pair. Without
    `pairs`, the inTime/outTime columns are the one pair."""
    if not p.pairs:
        return [("inTime", "IN", None), ("outTime", "OUT", None)]
    names = [str(h).strip() for h in header]
    out = []
    for pair in p.pairs:
        for key, direction in (("in", "IN"), ("out", "OUT")):
            src = str(pair.get(key) or "").strip()
            if src not in names:
                raise ParseRefused(f"The header has no column {src!r} for a declared pair.")
            i = names.index(src)
            out.append((src, direction, lambda cells, i=i: cells[i] if i < len(cells) else ""))
    return out


def _cell(index, target, row):
    f = index.get(target)
    v = f(row) if f else ""
    return v.strip() if isinstance(v, str) else v


# ------------------------------------------------------------------ parse


def parse_file(raw, p):
    """§4.1. Every error names its row, its field and quotes the raw value."""
    if len(raw) > MAX_BYTES:
        raise ParseRefused("The file is over 25 MB. Split the export by date range.")
    tz = zone(p.timezone)
    rows = list(rows_of(raw, p))
    header = rows[p.header_rows - 1][0] if p.header_rows and len(rows) >= p.header_rows else []
    index = _indexer(p, header if p.format != "FIXED" else [])
    if p.format != "FIXED" and p.header_rows:
        missing = [c["source"] for c in p.column_map if c.get("target") in TARGETS and index.get(c["target"]) is None]
        if missing:
            raise ParseRefused(f"The header has no column {', '.join(repr(m) for m in missing)}. Found: "
                               f"{', '.join(str(h) for h in header)}.")
    body = rows[p.header_rows:]
    pair_readers = _pairs(p, header)
    punches, errors, device_ids = [], [], {}

    def error(row, fld, value, message):
        errors.append({"row": row, "field": fld, "value": "" if value is None else str(value), "message": message})

    for n, (cells, raw_text) in enumerate(body):
        row = n + p.header_rows + 1
        ident = _cell(index, "deviceUserId", cells) or _cell(index, "employeeCode", cells)
        ident = str(ident).strip() if ident not in (None, "") else ""
        if p.combined_datetime:
            value = _cell(index, "datetime", cells)
            if isinstance(value, datetime):
                d, t = value.date(), value.time().replace(microsecond=0)
            else:
                parts = re.split(r"[ T]+", str(value or "").strip(), maxsplit=1)
                d = parse_date(parts[0], p.date_format, p.century_base)
                t = parse_time(parts[1], p.time_format) if len(parts) == 2 else None
            if d is None or t is None:
                error(row, "datetime", value, f'date and time "{value}" is not {p.date_format} {p.time_format}')
                continue
            times = [(t, _cell(index, "direction", cells))]
        else:
            raw_date = _cell(index, "date", cells)
            d = parse_date(raw_date, p.date_format, p.century_base)
            if d is None:
                error(row, "date", raw_date, f'date "{raw_date}" is not a valid date in {p.date_format}')
                continue
            times = None
        if not ident:
            error(row, "id", "", "no device ID or employee code in this row")
            continue
        terminal = str(_cell(index, "terminal", cells) or "")
        made = []
        if p.row_structure == "PAIRED_COLUMNS":
            for target, direction, read in pair_readers:
                value = read(cells) if read else _cell(index, target, cells)
                if value in (None, ""):
                    continue  # a missing out-punch is a fact for the register, not a parse error
                t = parse_time(value, p.time_format)
                if t is None:
                    error(row, target, value, f'time "{value}" is not a valid time in {p.time_format}')
                    continue
                made.append((t, direction, "PAIRED"))
            if not made:
                error(row, "inTime", "", "neither an in-time nor an out-time in this row")
                continue
        else:
            if times is None:
                value = _cell(index, "time", cells)
                t = parse_time(value, p.time_format)
                if t is None:
                    error(row, "time", value, f'time "{value}" is not a valid time in {p.time_format}')
                    continue
                times = [(t, _cell(index, "direction", cells))]
            t, said = times[0]
            if p.direction_handling == "EXPLICIT_COLUMN":
                made.append((t, "OUT" if OUT_WORDS.search(str(said or "").upper()) else "IN", "EXPLICIT"))
            else:
                made.append((t, None, p.direction_handling))
        for t, direction, source in made:
            instant, ld, lt = localise(d, t, tz)
            punches.append(Punch(row=row, raw_row=raw_text, device_user_id=ident, punched_at=instant, punch_date=ld,
                                 punch_time=lt, direction=direction, direction_source=source, terminal=terminal))
        device_ids[ident] = device_ids.get(ident, 0) + len(made)

    discarded = infer_directions(punches, p.direction_handling)
    dates = sorted(x.punch_date for x in punches)
    parsed_rows = len({x.row for x in punches})
    return Result(punches=punches, errors=errors, rows_total=len(body), rows_parsed=parsed_rows,
                  date_range=(dates[0], dates[-1]) if dates else None, device_ids=device_ids, discarded=discarded)


def infer_directions(punches, mode):
    """§4.3: group by employee and LOCAL day, then order -- alternating
    across the whole file starts the second employee of the day on OUT.
    FIRST_LAST keeps the first and last; the middle ones are marked
    discarded and stored as skipped, counted."""
    if mode not in ("ALTERNATE", "FIRST_LAST"):
        return 0
    by_day = {}
    for x in punches:
        by_day.setdefault((x.device_user_id, x.punch_date), []).append(x)
    discarded = 0
    for day in by_day.values():
        day.sort(key=lambda x: (x.punch_time, x.row))
        for i, x in enumerate(day):
            if mode == "ALTERNATE":
                x.direction = "IN" if i % 2 == 0 else "OUT"
            elif i == 0:
                x.direction = "IN"
            elif i == len(day) - 1:
                x.direction = "OUT"
            else:
                x.direction, x.discarded = None, True
                discarded += 1
    return discarded


def summary(result, limit=200):
    """The parse as the API shows it."""
    return {
        "rowsTotal": result.rows_total, "rowsParsed": result.rows_parsed, "rowsFailed": len({e["row"]
                                                                                            for e in result.errors}),
        "punches": len(result.punches), "discarded": result.discarded,
        "dateRange": {"from": result.date_range[0].isoformat(), "to": result.date_range[1].isoformat()}
        if result.date_range else None,
        "errors": result.errors[:limit], "errorsTotal": len(result.errors),
        "deviceIds": [{"deviceUserId": k, "punches": v} for k, v in sorted(result.device_ids.items())],
    }


def shift_minutes(t):
    return t.hour * 60 + t.minute


__all__ = ["Profile", "Punch", "Result", "ParseRefused", "parse_file", "parse_date", "parse_time", "timedelta"]
